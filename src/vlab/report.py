"""Les livrables : trois figures, un classeur Excel à formules vivantes, un mémo bilingue.

Style des figures commun au portfolio : palette d'Okabe et Ito, axes étiquetés, virgule décimale,
200 points par pouce. Le classeur n'est pas un export mort : la feuille DCF contient les FORMULES,
changer une hypothèse en haut recalcule la valeur par action en bas.
"""

from __future__ import annotations

from pathlib import Path

import matplotlib.pyplot as plt
import pandas as pd

from vlab.model import Hypotheses, fcff_path

OKABE_ITO = ["#0072B2", "#E69F00", "#009E73", "#D55E00", "#CC79A7", "#56B4E9", "#F0E442", "#000000"]


def use_style():
    import matplotlib as mpl
    from cycler import cycler
    from matplotlib.ticker import FuncFormatter

    mpl.rcParams.update({
        "figure.dpi": 200, "savefig.dpi": 200, "figure.constrained_layout.use": True,
        "font.size": 11, "axes.titlesize": 12, "axes.prop_cycle": cycler(color=OKABE_ITO),
        "axes.spines.top": False, "axes.spines.right": False,
        "axes.grid": True, "grid.alpha": 0.3, "grid.linewidth": 0.5,
        "legend.frameon": False, "lines.linewidth": 1.8,
    })
    return FuncFormatter(lambda v, _: f"{v:g}".replace(".", ","))


def fig_historique(h: pd.DataFrame, dest: Path) -> None:
    """Quinze ans de revenus (barres), marge d'exploitation et conversion en FCF (lignes)."""
    fr = use_style()
    fig, ax = plt.subplots(figsize=(9.5, 4.6))
    years = h.index.to_numpy()
    ax.bar(years, h["revenus"] / 1e3, color=OKABE_ITO[0], alpha=0.85, label="Revenus (G$ CAD, échelle gauche)")
    ax.set_ylabel("Revenus (milliards de CAD)")
    ax.yaxis.set_major_formatter(fr)
    ax2 = ax.twinx()
    ax2.plot(years, h["marge_exploitation_pct"], color=OKABE_ITO[3], marker="o", markersize=4,
             label="Marge d'exploitation (%, échelle droite)")
    ax2.plot(years, 100.0 * h["fcf"] / h["revenus"], color=OKABE_ITO[2], marker="s", markersize=4,
             label="FCF / revenus (%, échelle droite)")
    ax2.set_ylabel("En pourcent des revenus")
    ax2.yaxis.set_major_formatter(fr)
    ax2.spines["right"].set_visible(True)
    ax2.grid(False)
    if 2021 not in h.index:
        ax.annotate("2021 : absent du\nXBRL de la SEC", (2021, 8.0), ha="center", fontsize=8, color="0.35")
    ax.set_title("Le CN en quinze ans : des revenus réguliers, des marges de quasi-monopole")
    lines = ax.get_legend_handles_labels()[0] + ax2.get_legend_handles_labels()[0]
    labels = ax.get_legend_handles_labels()[1] + ax2.get_legend_handles_labels()[1]
    ax.legend(lines, labels, loc="lower right", fontsize=8.5)
    fig.savefig(dest)
    plt.close(fig)


def fig_football(price: float, dcf_range: tuple[float, float], ev_ebitda: float, pe: float,
                 base_dcf: float, dest: Path) -> None:
    """Les fourchettes de valorisation face au cours, la lecture d'un coup d'œil."""
    fr = use_style()
    fig, ax = plt.subplots(figsize=(9, 3.6))
    bars = [("DCF (sensibilité WACC x g)", dcf_range, OKABE_ITO[0]),
            ("Comparables EV/EBITDA (médiane des pairs)", (ev_ebitda * 0.95, ev_ebitda * 1.05), OKABE_ITO[2]),
            ("Comparables P/E (médiane des pairs)", (pe * 0.95, pe * 1.05), OKABE_ITO[4])]
    for i, (label, (lo, hi), color) in enumerate(bars):
        ax.barh(i, hi - lo, left=lo, height=0.5, color=color, alpha=0.8)
        ax.text(lo, i + 0.33, label, fontsize=9)
    ax.plot([base_dcf], [0], "D", color="black", markersize=6, zorder=5)
    ax.axvline(price, color=OKABE_ITO[3], linewidth=1.6)
    ax.text(price, len(bars) - 0.3, f" cours : {price:.0f} $", color=OKABE_ITO[3], fontsize=10)
    ax.set_yticks([])
    ax.set_xlabel("Valeur par action (CAD)")
    ax.xaxis.set_major_formatter(fr)
    ax.set_title("Trois approches, un cours : où tombe le marché")
    fig.savefig(dest)
    plt.close(fig)


def fig_croissance_implicite(implied: float, hist_fcff: float, hist_rev: float, dest: Path) -> None:
    """La croissance que le cours suppose, face à ce que l'entreprise a réellement fait."""
    fr = use_style()
    fig, ax = plt.subplots(figsize=(8, 3.2))
    labels = ["Croissance implicite dans le cours\n(DCF inversé, 5 premières années)",
              "FCFF réalisé, croissance annuelle\nmoyenne des 10 derniers exercices",
              "Revenus réalisés, croissance annuelle\nmoyenne des 10 derniers exercices"]
    vals = [100 * implied, 100 * hist_fcff, 100 * hist_rev]
    colors = [OKABE_ITO[3], OKABE_ITO[0], OKABE_ITO[2]]
    bars = ax.barh(labels[::-1], vals[::-1], color=colors[::-1], height=0.55)
    for rect, v in zip(bars, vals[::-1], strict=True):
        ax.text(rect.get_width() + 0.08, rect.get_y() + rect.get_height() / 2,
                f"{v:.1f} %".replace(".", ","), va="center", fontsize=10)
    ax.set_xlabel("Croissance annuelle (%)")
    ax.xaxis.set_major_formatter(fr)
    ax.set_title("Ce que le cours suppose contre ce que le CN a livré")
    fig.savefig(dest)
    plt.close(fig)


def build_workbook(hist: pd.DataFrame, h: Hypotheses, sens: pd.DataFrame, comps: pd.DataFrame,
                   market_date: str, dest: Path) -> None:
    """Le classeur : Lisez-moi, Historique (valeurs), DCF (FORMULES), Sensibilité, Comparables."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    bold = Font(bold=True)

    ws = wb.active
    ws.title = "Lisez-moi"
    ws["A1"] = "Valorisation du Canadien National (CNR.TO), classeur généré par vlab"
    ws["A1"].font = bold
    ws["A3"] = "La feuille DCF contient des FORMULES : changer une hypothèse (en haut) recalcule tout."
    ws["A4"] = "Statuts : historique MESURÉ (SEC, companyfacts) ; hypothèses DCF MODÉLISÉES et déclarées ;"
    ws["A5"] = f"multiples des pairs RAPPORTÉS (Yahoo Finance, {market_date}). Montants en millions de CAD."

    ws = wb.create_sheet("Historique")
    ws.append(["exercice", *hist.columns.tolist()])
    for c in ws[1]:
        c.font = bold
    for year, row in hist.iterrows():
        ws.append([int(year), *[None if pd.isna(v) else round(float(v), 3) for v in row]])

    ws = wb.create_sheet("DCF")
    ws["A1"], ws["A1"].font = "Hypothèses (modifiables)", bold
    inputs = [("FCFF de départ (M$ CAD)", h.fcff_base), ("WACC", h.wacc),
              ("Croissance années 1 à 5", h.g_initial), ("Croissance perpétuelle", h.g_terminal),
              ("Dette nette (M$ CAD)", h.dette_nette), ("Actions diluées (millions)", h.actions)]
    for i, (label, value) in enumerate(inputs, start=2):
        ws[f"A{i}"], ws[f"B{i}"] = label, round(float(value), 4)
    ws["A9"], ws["A9"].font = "Projection", bold
    ws.append([])
    ws["A10"], ws["B10"], ws["C10"], ws["D10"] = "année", "croissance", "FCFF (M$)", "valeur actuelle (M$)"
    for c in ws[10]:
        c.font = bold
    for t in range(1, h.annees + 1):
        r = 10 + t
        ws[f"A{r}"] = t
        if t <= 5:
            ws[f"B{r}"] = "=$B$4"
        else:
            ws[f"B{r}"] = f"=(1-({t}-5)/5)*$B$4+(({t}-5)/5)*$B$5"
        prev = "$B$2" if t == 1 else f"C{r - 1}"
        ws[f"C{r}"] = f"={prev}*(1+B{r})"
        ws[f"D{r}"] = f"=C{r}/(1+$B$3)^A{r}"
    last = 10 + h.annees
    ws[f"A{last + 2}"] = "Valeur terminale (Gordon), actualisée"
    ws[f"D{last + 2}"] = f"=C{last}*(1+$B$5)/($B$3-$B$5)/(1+$B$3)^{h.annees}"
    ws[f"A{last + 3}"] = "Valeur d'entreprise (M$)"
    ws[f"D{last + 3}"] = f"=SUM(D11:D{last})+D{last + 2}"
    ws[f"A{last + 4}"] = "Valeur par action (CAD)"
    ws[f"D{last + 4}"] = f"=(D{last + 3}-$B$6)/$B$7"
    ws[f"A{last + 4}"].font = bold
    ws[f"D{last + 4}"].font = bold

    ws = wb.create_sheet("Sensibilite")
    ws["A1"] = "Valeur par action (CAD) selon WACC (lignes) et croissance perpétuelle (colonnes)"
    ws["A1"].font = bold
    ws.append(["wacc \\ g_terminal", *[f"{g:.1%}" for g in sens.columns]])
    for w, row in sens.iterrows():
        ws.append([f"{w:.1%}", *[None if pd.isna(v) else round(float(v), 1) for v in row]])

    ws = wb.create_sheet("Comparables")
    ws.append(["titre", "EV/EBITDA", "P/E"])
    for c in ws[1]:
        c.font = bold
    for t, row in comps.iterrows():
        ws.append([t, *[None if pd.isna(v) else round(float(v), 2) for v in row]])

    for sheet in wb.worksheets:
        for col in ("A", "B", "C", "D"):
            sheet.column_dimensions[col].width = max(sheet.column_dimensions[col].width or 0, 26 if col == "A" else 14)
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)


def project_flows_for_memo(h: Hypotheses) -> float:
    return float(fcff_path(h)[0])
